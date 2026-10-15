"""Demo workbook: ~300 synthetic players for the UltraByte import (Add mode). Deterministic."""
import random, sys
from datetime import datetime, timedelta
from openpyxl import Workbook
sys.path.insert(0, '.')
from app.importer import SHEETS

rng = random.Random(20261015)
CITIES = ["Toronto","Ottawa","Hamilton","London","Kitchener","Windsor","Sudbury","Sault Ste. Marie","Thunder Bay","Kingston","Barrie","Niagara Falls"]
AGE = ["19-29","30-44","45-64","65+"]; INC = ["<20k","20-50k","50-100k","100-200k","200k+"]
OCC = {"<20k":["student","part-time server","retail associate"],"20-50k":["cook","warehouse worker","admin assistant","retired"],
       "50-100k":["nurse","teacher","electrician","analyst"],"100-200k":["engineer","sales manager","pharmacist"],"200k+":["business owner","physician","lawyer"]}
MONTHLY = {"<20k":120,"20-50k":250,"50-100k":500,"100-200k":1200,"200k+":3000}
FIRST = ["Ava","Liam","Noah","Emma","Mia","Ethan","Zoe","Lucas","Maya","Owen","Leah","Kai","Nora","Eli","Iris","Theo","Ruby","Finn","Ada","Jude"]
H0 = datetime(2026,4,15); W0 = datetime(2026,9,15); W1 = datetime(2026,10,14,23)
players, txns, sessions = [], [], []
n = 0
def pid():
    global n; n += 1; return f"D-{n:03d}"
def name(): return f"{rng.choice(FIRST)} {rng.choice('ABCDEFGHJKLMNPRSTVW')}."
def add_player(inc=None, age=None, city=None, occ=None):
    inc = inc or rng.choice(INC); age = age or rng.choice(AGE); city = city or rng.choice(CITIES)
    p = dict(id=pid(), name=name(), occ=occ or rng.choice(OCC[inc]), inc=inc, age=age, city=city)
    players.append(p); return p
def ts(d, h): return (d + timedelta(hours=h)).strftime("%Y-%m-%d %H:%M")
def dep(p, when, amt, method="debit", card=None, last4=None, city=None, dev=None):
    txns.append([p["id"], when, "deposit", round(amt,2), method, card or "", last4 or "", city or p["city"], dev or f"DEV-{p['id']}", "", ""])
def wd(p, when, amt, acct=None, match="yes", city=None):
    txns.append([p["id"], when, "withdrawal", round(amt,2), "e_transfer", "", "", city or p["city"], f"DEV-{p['id']}", acct or f"ACCT-{p['id']}", match])
def wager(p, when, amt, city=None): txns.append([p["id"], when, "wager", round(amt,2), "wallet", "", "", city or p["city"], f"DEV-{p['id']}", "", ""])
def sess(p, start, hours, net, city=None, dev=None):
    sessions.append([p["id"], ts(start,0), ts(start, hours), round(net,2), city or p["city"], dev or f"DEV-{p['id']}"])
def normal_history(p, monthly, months=6, night=False, rising=False):
    for k in range(months):
        base = H0 + timedelta(days=30*k)
        m = monthly * (1 + 0.6*k if rising else rng.uniform(0.85,1.15))
        cnt = min(8, max(2, round(m/80)))
        days = rng.sample(range(30), cnt)
        for day in days:
            d = base + timedelta(days=day); h = rng.uniform(0.2,3.5) if (night and rng.random()<0.85) else rng.uniform(10,21)
            amt = m/cnt*rng.uniform(0.6,1.4); dep(p, ts(d,h), amt); sess(p, d+timedelta(hours=h+0.1), rng.uniform(0.5,2.5), -amt*rng.uniform(0.2,0.9)); wager(p, ts(d,h+0.2), amt*1.3)
# --- 265 ordinary players
for _ in range(265):
    p = add_player(); normal_history(p, MONTHLY[p["inc"]])
# --- 10 structuring (STR): 2-4 deposits 9,000-9,999 in the window
for _ in range(10):
    p = add_player(inc=rng.choice(["<20k","20-50k","50-100k"])); normal_history(p, MONTHLY[p["inc"]], months=5)
    for j in range(rng.randint(2,4)):
        d = W0 + timedelta(days=rng.randrange(25)); dep(p, ts(d,14+j), rng.uniform(9000,9990), method="cash_cage"); sess(p, d+timedelta(hours=15), 1, -300); wager(p, ts(d,15.1), 400)
# --- 6 account takeover (URGENT): 3 cities in a day + new card + new payout account
for _ in range(6):
    p = add_player(); normal_history(p, MONTHLY[p["inc"]], months=5)
    d = W0 + timedelta(days=rng.randrange(20)); cities = rng.sample([c for c in CITIES if c != p["city"]], 2)
    sess(p, d+timedelta(hours=8), 0.3, -40); sess(p, d+timedelta(hours=12), 0.3, -30, city=cities[0], dev="DEV-NEW"); sess(p, d+timedelta(hours=16), 0.3, -20, city=cities[1], dev="DEV-NEW")
    dep(p, ts(d,16.2), 900, method="credit", card=f"CARD-NEW-{p['id']}", last4=str(rng.randint(1000,9999)), city=cities[1], dev="DEV-NEW")
    wd(p, ts(d,16.8), 850, acct=f"ACCT-NEW-{p['id']}", match="yes", city=cities[1])
# --- 5 minimal play then withdraw (STR)
for _ in range(5):
    p = add_player(); normal_history(p, MONTHLY[p["inc"]], months=5)
    d = W0 + timedelta(days=rng.randrange(20)); amt = rng.uniform(6000,15000)
    dep(p, ts(d,18), amt, method="e_transfer"); wager(p, ts(d,18.5), amt*0.03); sess(p, d+timedelta(hours=18.4), 0.4, -amt*0.03); wd(p, ts(d+timedelta(days=1),10), amt*0.95)
# --- 4 third-party payout (STR)
for _ in range(4):
    p = add_player(); normal_history(p, MONTHLY[p["inc"]])
    d = W0 + timedelta(days=rng.randrange(25)); wd(p, ts(d,11), rng.uniform(1500,5000), acct=f"ACCT-OTHER-{p['id']}", match="no")
# --- 3 pairs sharing a card (STR)
for k in range(3):
    a, b = add_player(), add_player(); normal_history(a, MONTHLY[a["inc"]]); normal_history(b, MONTHLY[b["inc"]])
    for p in (a, b):
        d = W0 + timedelta(days=rng.randrange(25)); dep(p, ts(d,13), rng.uniform(200,600), card=f"CARD-SHARED-{k}", last4=f"55{k}{k}")
# --- 5 Safer Gambling: rising deposits + night play (escalation + night)
for _ in range(5):
    p = add_player(); normal_history(p, MONTHLY[p["inc"]]*0.6, rising=True, night=True)
# --- 4 loss chasing + intensity: deposits right after losing sessions, most days
for _ in range(4):
    p = add_player(); normal_history(p, MONTHLY[p["inc"]], months=5)
    for day in range(0, 28):
        if rng.random() < 0.8:
            d = W0 + timedelta(days=day); s0 = d + timedelta(hours=20); sess(p, s0, 1.0, -60); dep(p, ts(d, 21.3), 60); wager(p, ts(d,21.4), 70)
# --- 3 controls: single large deposit, consistent with income (should NOT alert in V2, alerts in V1)
for _ in range(3):
    p = add_player(inc="200k+", occ="business owner"); normal_history(p, 3000)
    d = W0 + timedelta(days=rng.randrange(25)); dep(p, ts(d,19), 10500, method="debit"); sess(p, d+timedelta(hours=19.2), 3, -3000); wager(p, ts(d,19.3), 7000)
wb = Workbook(); wb.remove(wb.active)
ws = wb.create_sheet("players"); ws.append(SHEETS["players"])
# One display name per player, none shared with the built-in seed (team feedback 2026-10-15). Fixed order, no random
# draws, so every transaction and score is unchanged.
from db.seed.generate import FIRST as SEED_FIRST
_combos = [f"{f} {l}." for f in FIRST if f not in SEED_FIRST for l in "ABCDEFGHJKLMNPRSTVW"]
_used = set()
for p in players:
    if p["name"] in _used or p["name"].split()[0] in SEED_FIRST:
        p["name"] = next(c for c in _combos if c not in _used)
    _used.add(p["name"])
for p in players: ws.append([p["id"], p["name"], p["occ"], p["inc"], p["age"], p["city"]])
ws = wb.create_sheet("transactions"); ws.append(SHEETS["transactions"])
for t in sorted(txns, key=lambda t: t[1]): ws.append(t)
ws = wb.create_sheet("sessions"); ws.append(SHEETS["sessions"])
for s in sorted(sessions, key=lambda s: s[1]): ws.append(s)
ws = wb.create_sheet("README"); 
for line in ["UltraByte demo workbook. ALL DATA IS SYNTHETIC (generated 2026-10-15 for ThunderHacks II).",
             f"{len(players)} players, {len(txns)} transactions, {len(sessions)} sessions. Player ids D-001..D-{len(players):03d}.",
             "Import with mode ADD (keeps the 490 background players so the peer check has enough comparable players).",
             "Scenarios (ids in order): D-001..265 ordinary; 266..275 structuring (STR); 276..281 account takeover (URGENT); 282..286 deposit, little play, withdraw (STR); 287..290 third-party payout (STR); 291..296 three shared-card pairs (STR); 297..301 rising deposits + night play (Safer Gambling); 302..305 loss chasing + intensity (Safer Gambling); 306..308 single $10,500 deposit by a high earner: Version 1.0 alerts, Version 2.0 gives context (Low/Medium, no label).",
             "To remove: Review queue -> Reset demo data (restores the synthetic seed)."]: ws.append([line])
wb.save("demo/UltraByte_demo_300_players.xlsx")
print(len(players), len(txns), len(sessions))
