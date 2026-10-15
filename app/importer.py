"""Import synthetic data from an Excel workbook or CSV files (owner decision 2026-10-14 20:40 EDT, HS4).

Three sheets: players, transactions, sessions. Every row is validated before anything is loaded;
a file with errors loads nothing. Sensitive columns (gender, ethnicity, exact age, names of
account holders) are rejected on purpose: data minimization is part of the product.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook

from app.scoring import config as C

AGE_BANDS = ["19-29", "30-44", "45-64", "65+"]
INCOME_BANDS = list(C.INCOME_MIDPOINT)
TXN_TYPES = ["deposit", "withdrawal", "withdrawal_reversal", "wager", "win"]
METHODS = ["debit", "credit", "e_transfer", "cash_cage", "prepaid", "wallet"]
FORBIDDEN = {"gender", "sex", "ethnicity", "race", "religion", "age", "date_of_birth", "dob", "birthdate",
             "account_holder_name", "holder_name", "card_number", "account_number", "sin", "ssn", "email", "phone"}

SHEETS = {
    "players": ["player_id", "display_name", "occupation", "income_band", "age_band", "home_city"],
    "transactions": ["player_id", "timestamp", "type", "amount", "method", "card_ref", "card_last4", "city",
                     "device_ref", "payout_account_ref", "payout_holder_matches"],
    "sessions": ["player_id", "start", "end", "net_result", "city", "device_ref"],
}
REQUIRED = {
    "players": ["player_id", "display_name", "occupation", "income_band", "age_band", "home_city"],
    "transactions": ["player_id", "timestamp", "type", "amount", "method", "city"],
    "sessions": ["player_id", "start", "end", "net_result", "city"],
}
HELP = {
    "players": "One row per player. Bands only: no exact age, gender or ethnicity (such columns are rejected).",
    "transactions": "Deposits, withdrawals, wagers, wins. Timestamps between 2026-04-15 and 2026-10-14 (EDT). "
                    "card_ref / payout_account_ref are any stable labels (e.g. CARD-A); use the same label on two "
                    "players to simulate a shared card or account. payout_holder_matches: yes/no on withdrawals.",
    "sessions": "One row per play session, with start and end timestamps and net result (wins minus wagers).",
}
EXAMPLES = {
    "players": [["P-0001", "Sam K.", "student", "<20k", "19-29", "Sault Ste. Marie"],
                ["P-0002", "Lee M.", "business owner", "200k+", "45-64", "Toronto"]],
    "transactions": [["P-0001", "2026-10-01 14:00", "deposit", 9500, "cash_cage", "CARD-A", "4417", "Sault Ste. Marie", "DEV-1", "", ""],
                     ["P-0001", "2026-10-02 10:00", "withdrawal", 4000, "e_transfer", "", "", "Sault Ste. Marie", "DEV-1", "ACCT-X", "no"],
                     ["P-0002", "2026-10-03 19:00", "deposit", 10500, "debit", "CARD-B", "1234", "Toronto", "DEV-2", "", ""],
                     ["P-0002", "2026-10-03 19:30", "wager", 6000, "wallet", "", "", "Toronto", "DEV-2", "", ""]],
    "sessions": [["P-0001", "2026-10-01 14:10", "2026-10-01 16:00", -800, "Sault Ste. Marie", "DEV-1"],
                 ["P-0002", "2026-10-03 19:10", "2026-10-03 22:00", -2400, "Toronto", "DEV-2"]],
}


def template_xlsx() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, cols in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for row in EXAMPLES[name]:
            ws.append(row)
    ws = wb.create_sheet("README")
    ws.append(["UltraByte import template. All data must be synthetic. Delete the example rows or keep them."])
    for name in SHEETS:
        ws.append([f"{name}: {HELP[name]}"])
    ws.append([f"income_band: {', '.join(INCOME_BANDS)}  |  age_band: {', '.join(AGE_BANDS)}"])
    ws.append([f"transaction type: {', '.join(TXN_TYPES)}  |  method: {', '.join(METHODS)}"])
    ws.append(["Windows: history 2026-04-15 to 2026-09-14, current window 2026-09-15 to 2026-10-14 (EDT)."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass
class Result:
    players: list = field(default_factory=list)
    txns: list = field(default_factory=list)
    sessions: list = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def _read_tables(filename: str, data: bytes) -> dict[str, list[dict]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = {}
        for sheet in SHEETS:
            if sheet not in wb.sheetnames:
                continue
            rows = list(wb[sheet].iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
            out[sheet] = [dict(zip(header, r)) for r in rows[1:] if any(v is not None and str(v).strip() != "" for v in r)]
        return out
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig")
        rows = list(csv.DictReader(io.StringIO(text)))
        header = {k.strip().lower() for k in (rows[0].keys() if rows else [])}
        for sheet, cols in SHEETS.items():
            if set(REQUIRED[sheet]) <= header:
                return {sheet: [{k.strip().lower(): v for k, v in r.items()} for r in rows]}
        raise ValueError("CSV header does not match any sheet (players, transactions or sessions)")
    raise ValueError("Upload an .xlsx workbook or a .csv file")


def _ts(v, where: str, errors: list[str]):
    if isinstance(v, datetime):
        dt = v
    else:
        s = str(v).strip().replace("T", " ")
        dt = None
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
            try:
                dt = datetime.strptime(s, fmt)
                break
            except ValueError:
                pass
        if dt is None:
            errors.append(f"{where}: cannot read timestamp '{v}' (use YYYY-MM-DD HH:MM)")
            return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=C.EDT)
    if not (C.HIST_START <= dt < C.WIN_END):
        errors.append(f"{where}: timestamp {dt:%Y-%m-%d} is outside 2026-04-15 to 2026-10-14")
        return None
    return dt


def _money(v, where, errors):
    try:
        x = float(v)
    except (TypeError, ValueError):
        errors.append(f"{where}: amount '{v}' is not a number")
        return None
    return round(x, 2)


def _yn(v):
    if v is None or str(v).strip() == "":
        return None
    return str(v).strip().lower() in ("yes", "y", "true", "1", "match")


def parse(filename: str, data: bytes) -> Result:
    res = Result()
    try:
        tables = _read_tables(filename, data)
    except Exception as exc:
        res.errors.append(str(exc))
        return res
    if "players" not in tables:
        res.errors.append("Sheet 'players' is required")
        return res

    for sheet, rows in tables.items():
        header = set(rows[0].keys()) if rows else set()
        bad = sorted(header & FORBIDDEN)
        if bad:
            res.errors.append(f"{sheet}: columns not allowed (data minimization): {', '.join(bad)}")
        missing = [c for c in REQUIRED[sheet] if c not in header]
        if rows and missing:
            res.errors.append(f"{sheet}: missing columns: {', '.join(missing)}")
    if res.errors:
        return res

    ids = set()
    for i, r in enumerate(tables["players"], start=2):
        w = f"players row {i}"
        pid = str(r.get("player_id") or "").strip()
        if not pid:
            res.errors.append(f"{w}: player_id is empty"); continue
        if pid in ids:
            res.errors.append(f"{w}: duplicate player_id {pid}"); continue
        inc, age = str(r.get("income_band") or "").strip(), str(r.get("age_band") or "").strip()
        if inc not in INCOME_BANDS:
            res.errors.append(f"{w}: income_band '{inc}' must be one of {', '.join(INCOME_BANDS)}")
        if age not in AGE_BANDS:
            res.errors.append(f"{w}: age_band '{age}' must be one of {', '.join(AGE_BANDS)}")
        ids.add(pid)
        res.players.append(dict(id=pid, display_name=str(r.get("display_name") or pid)[:60],
                                occupation=str(r.get("occupation") or "unknown")[:60], income_band=inc, age_band=age,
                                region=str(r.get("home_city") or "unknown")[:60], peer_segment=f"{age}|{inc}",
                                created_at=C.HIST_START - timedelta(days=60)))

    for i, r in enumerate(tables.get("transactions", []), start=2):
        w = f"transactions row {i}"
        pid = str(r.get("player_id") or "").strip()
        if pid not in ids:
            res.errors.append(f"{w}: unknown player_id {pid}"); continue
        ts = _ts(r.get("timestamp"), w, res.errors)
        typ = str(r.get("type") or "").strip().lower()
        if typ not in TXN_TYPES:
            res.errors.append(f"{w}: type '{typ}' must be one of {', '.join(TXN_TYPES)}")
        amt = _money(r.get("amount"), w, res.errors)
        if amt is not None and amt <= 0:
            res.errors.append(f"{w}: amount must be positive")
        method = str(r.get("method") or "wallet").strip().lower()
        if method not in METHODS:
            res.errors.append(f"{w}: method '{method}' must be one of {', '.join(METHODS)}")
        if ts is None or amt is None:
            continue
        card = str(r.get("card_ref") or "").strip() or None
        acct = str(r.get("payout_account_ref") or "").strip() or None
        res.txns.append(dict(player_id=pid, ts=ts, type=typ, amount=amt, method=method,
                             method_fingerprint=("imp-" + card) if card and typ == "deposit" else None,
                             method_last4=str(r.get("card_last4") or "")[-4:] or None if typ == "deposit" else None,
                             location_region=str(r.get("city") or "unknown")[:60],
                             device_id=("imp-" + str(r.get("device_ref")).strip()) if r.get("device_ref") else None,
                             payout_account=("imp-" + acct) if acct and typ == "withdrawal" else None,
                             payout_holder_match=_yn(r.get("payout_holder_matches")) if typ == "withdrawal" else None))

    for i, r in enumerate(tables.get("sessions", []), start=2):
        w = f"sessions row {i}"
        pid = str(r.get("player_id") or "").strip()
        if pid not in ids:
            res.errors.append(f"{w}: unknown player_id {pid}"); continue
        a, b = _ts(r.get("start"), w, res.errors), _ts(r.get("end"), w, res.errors)
        net = _money(r.get("net_result"), w, res.errors)
        if a is None or b is None or net is None:
            continue
        if b < a:
            res.errors.append(f"{w}: end is before start"); continue
        res.sessions.append(dict(player_id=pid, start_ts=a, end_ts=b, net_result=net,
                                 location_region=str(r.get("city") or "unknown")[:60],
                                 device_id=("imp-" + str(r.get("device_ref")).strip()) if r.get("device_ref") else None))

    if not res.txns:
        res.warnings.append("No transactions: every imported player will score 0.")
    return res


def load(conn, res: Result, mode: str) -> dict:
    """mode 'replace' = only the imported data; 'append' = add to what is loaded now."""
    from db.seed.generate import SES_COLS, TXN_COLS

    with conn.cursor() as cur:
        if mode == "replace":
            # Like /reset: the demo-only path that clears the audit trail (audit rows reference players).
            cur.execute("TRUNCATE audit_event, planted_case, txn, session, player RESTART IDENTITY")
        else:
            existing = {r["id"] for r in cur.execute("SELECT id FROM player").fetchall()}
            clash = sorted(p["id"] for p in res.players if p["id"] in existing)
            if clash:
                raise ValueError(f"player_id already exists: {', '.join(clash[:5])}{'...' if len(clash) > 5 else ''}")
        with cur.copy("COPY player (id, display_name, occupation, income_band, age_band, region, peer_segment, created_at, entry_source) FROM STDIN") as cp:
            for p in res.players:
                cp.write_row((p["id"], p["display_name"], p["occupation"], p["income_band"], p["age_band"], p["region"], p["peer_segment"], p["created_at"], "import"))
        with cur.copy(f"COPY txn ({', '.join(TXN_COLS)}) FROM STDIN") as cp:
            for t in sorted(res.txns, key=lambda t: t["ts"]):
                cp.write_row(tuple(t[c] for c in TXN_COLS))
        with cur.copy(f"COPY session ({', '.join(SES_COLS)}) FROM STDIN") as cp:
            for s in sorted(res.sessions, key=lambda s: s["start_ts"]):
                cp.write_row(tuple(s[c] for c in SES_COLS))
        cur.execute("INSERT INTO audit_event (actor, action, note) VALUES ('investigator (demo)', 'reset', %s)",
                    (f"Imported synthetic data ({mode}): {len(res.players)} players, {len(res.txns)} transactions, {len(res.sessions)} sessions",))
    conn.commit()
    return {"players": len(res.players), "txns": len(res.txns), "sessions": len(res.sessions)}
